"""تحليل آمن ومحدود لملفات CSV/XLSX المرفقة بالمحادثة.

الأداة لا تنفّذ أي كود يقدمه المستخدم. هي تقرأ البيانات فقط وتحسب
إحصاءات وصفية بسيطة يمكن تمريرها إلى المساعد أو عرضها مباشرة.
"""
from __future__ import annotations

import csv
import math
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook


MAX_FILENAME_LENGTH = 255
MAX_OUTPUT_CHARS = 12_000
MAX_SAMPLE_ROWS = 5
MAX_TOP_VALUES = 5


class DataAnalysisError(ValueError):
    """خطأ آمن ومفهوم للمستخدم أثناء تحليل البيانات."""


@dataclass(frozen=True)
class DataFile:
    path: Path
    original_filename: str
    content_type: str


def extract_data_analysis_request(message: str) -> str | None:
    """يرجع اسم الملف إذا كانت الرسالة أمر تحليل، وإلا None."""
    text = message.strip()
    for prefix in ("/analyze", "/analyse", "/data"):
        if text == prefix:
            return ""
        if text.startswith(prefix + " "):
            return text[len(prefix) + 1 :].strip()
    return None


def _clean_name(value: str) -> str:
    return value.strip()[:MAX_FILENAME_LENGTH]


def _to_number(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if not text:
            return None
        try:
            number = float(text)
        except ValueError:
            return None
        return number if math.isfinite(number) else None
    return None


def _format_number(value: float) -> str:
    if value.is_integer():
        return str(int(value))
    return f"{value:.4f}".rstrip("0").rstrip(".")


def _analyze_rows(headers: list[str], rows: list[list[object]], source_label: str) -> str:
    if not headers:
        raise DataAnalysisError("لم أجد أعمدة واضحة داخل الملف.")

    normalized_headers = []
    seen_headers: Counter[str] = Counter()
    for index, raw_header in enumerate(headers, start=1):
        name = str(raw_header).strip() if raw_header is not None else ""
        if not name:
            name = f"column_{index}"
        seen_headers[name] += 1
        if seen_headers[name] > 1:
            name = f"{name}_{seen_headers[name]}"
        normalized_headers.append(name)

    width = len(normalized_headers)
    normalized_rows = [
        list(row[:width]) + [None] * max(0, width - len(row))
        for row in rows
    ]

    lines = [
        f"## تحليل البيانات: {source_label}",
        "",
        f"- الصفوف: **{len(normalized_rows)}**",
        f"- الأعمدة: **{len(normalized_headers)}**",
        "",
        "### الأعمدة",
    ]

    for index, header in enumerate(normalized_headers):
        values = [row[index] for row in normalized_rows]
        non_empty = [value for value in values if value is not None and str(value).strip() != ""]
        missing = len(values) - len(non_empty)
        numbers = [number for value in non_empty if (number := _to_number(value)) is not None]

        line = f"- **{header}** — ممتلئ: {len(non_empty)}, فارغ: {missing}"
        if numbers and len(numbers) >= max(2, math.ceil(len(non_empty) * 0.6)):
            line += (
                f", متوسط: {_format_number(mean(numbers))},"
                f" أدنى: {_format_number(min(numbers))},"
                f" أقصى: {_format_number(max(numbers))}"
            )
        else:
            counter = Counter(str(value).strip() for value in non_empty)
            common = counter.most_common(MAX_TOP_VALUES)
            if common:
                top_values = ", ".join(
                    f"{value} ({count})" for value, count in common
                )
                line += f", الأكثر تكرارًا: {top_values}"
        lines.append(line)

    if normalized_rows:
        lines.extend(["", "### عينة من أول الصفوف"])
        sample_headers = " | ".join(normalized_headers)
        lines.append(f"| {sample_headers} |")
        lines.append("| " + " | ".join("---" for _ in normalized_headers) + " |")
        for row in normalized_rows[:MAX_SAMPLE_ROWS]:
            cells = [
                "" if value is None else str(value).replace("|", "\\|").replace("\n", " ")
                for value in row
            ]
            lines.append(f"| {' | '.join(cells)} |")

    result = "\n".join(lines).strip()
    if len(result) > MAX_OUTPUT_CHARS:
        return result[:MAX_OUTPUT_CHARS].rstrip() + "\n\n[تم اختصار تقرير التحليل]"
    return result


def analyze_csv(path: Path, original_filename: str) -> str:
    try:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
    except OSError as exc:
        raise DataAnalysisError("تعذر قراءة الملف.") from exc

    try:
        reader = csv.reader(text.splitlines())
        all_rows = list(reader)
    except csv.Error as exc:
        raise DataAnalysisError("ملف CSV غير صالح.") from exc

    if not all_rows:
        raise DataAnalysisError("الملف فارغ.")

    headers = all_rows[0]
    rows = [list(row) for row in all_rows[1:]]
    return _analyze_rows(headers, rows, original_filename)


def analyze_xlsx(path: Path, original_filename: str) -> str:
    try:
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise DataAnalysisError("تعذر فتح ملف Excel.") from exc

    sheets = []
    try:
        for worksheet in workbook.worksheets:
            values = list(worksheet.iter_rows(values_only=True))
            if not values:
                continue
            sheet_name = worksheet.title
            report = _analyze_rows(
                list(values[0]),
                [list(row) for row in values[1:]],
                f"{original_filename} — {sheet_name}",
            )
            sheets.append(report)
    finally:
        workbook.close()

    if not sheets:
        raise DataAnalysisError("لم أجد بيانات داخل ملف Excel.")
    combined = "\n\n".join(sheets)
    if len(combined) > MAX_OUTPUT_CHARS:
        combined = combined[:MAX_OUTPUT_CHARS].rstrip() + "\n\n[تم اختصار تقرير التحليل]"
    return combined


def analyze_file(file: DataFile) -> str:
    suffix = file.path.suffix.lower()
    if suffix == ".csv" or file.content_type == "text/csv":
        return analyze_csv(file.path, file.original_filename)
    if suffix == ".xlsx" or file.content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return analyze_xlsx(file.path, file.original_filename)
    raise DataAnalysisError("تحليل البيانات يدعم ملفات CSV وXLSX فقط حاليًا.")


__all__ = [
    "DataAnalysisError",
    "DataFile",
    "analyze_file",
    "extract_data_analysis_request",
]
