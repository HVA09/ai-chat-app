"""استخراج نص آمن ومحدود من الملفات القابلة للقراءة قبل إدخالها إلى سياق AI."""
from __future__ import annotations

import csv
from io import StringIO
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pypdf import PdfReader


SUPPORTED_TEXT_TYPES = {
    "text/plain",
    "text/csv",
    "application/vnd.ms-excel",
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

DEFAULT_MAX_CHARS = 20_000


class FileTextExtractionError(Exception):
    """فشل في قراءة محتوى الملف."""


def _limit(text: str, max_chars: int) -> str:
    text = text.strip()
    if len(text) <= max_chars:
        return text
    return f"{text[:max_chars]}\n\n[تم اختصار المحتوى بسبب الحد المسموح]"


def _read_text_file(path: Path, max_chars: int) -> str:
    # Path.read_bytes() loads the entire file before slicing. Read a bounded
    # byte prefix directly from the file handle so large uploads do not cause
    # avoidable peak-memory spikes before the character limit is applied.
    with path.open("rb") as source:
        raw = source.read(max(0, max_chars) * 4)
    text = raw.decode("utf-8-sig", errors="replace")
    return _limit(text, max_chars)


def _read_csv_file(path: Path, max_chars: int) -> str:
    # Bound the input bytes as well as the rendered output. A single very large
    # CSV field can otherwise be allocated in full before the row limit applies.
    with path.open("rb") as source:
        raw = source.read(max(0, max_chars) * 8)
    text = raw.decode("utf-8-sig", errors="replace")
    rows = csv.reader(StringIO(text))
    rendered = []
    used = 0
    for row in rows:
        line = " | ".join(cell.strip() for cell in row)
        if not line:
            continue
        remaining = max_chars - used
        if remaining <= 0:
            break
        separator = 1 if rendered else 0
        available = max(0, remaining - separator)
        if len(line) > available:
            if available:
                rendered.append(line[:available])
                used += available + separator
            break
        rendered.append(line)
        used += len(line) + separator
        if used >= max_chars:
            break
    return _limit("\n".join(rendered), max_chars)


def _read_pdf_file(path: Path, max_chars: int) -> str:
    reader = PdfReader(str(path))
    parts: list[str] = []
    used = 0
    truncated = False

    for page in reader.pages:
        page_text = (page.extract_text() or "").strip()
        if not page_text:
            continue

        separator = 2 if parts else 0
        remaining = max_chars - used - separator
        if remaining <= 0:
            truncated = True
            break

        if len(page_text) >= remaining:
            parts.append(page_text[:remaining])
            used += separator + remaining
            truncated = True
            break

        parts.append(page_text)
        used += separator + len(page_text)

    content = "\n\n".join(parts).strip()
    if truncated and content:
        return f"{content}\n\n[تم اختصار المحتوى بسبب الحد المسموح]"
    return _limit(content, max_chars)


def _read_docx_file(path: Path, max_chars: int) -> str:
    document = Document(str(path))
    parts: list[str] = []
    used = 0
    truncated = False

    def append_part(value: str) -> bool:
        """Append only the portion that fits; return True once text was clipped."""
        nonlocal used, truncated
        value = value.strip()
        if not value:
            return False

        separator = 1 if parts else 0
        remaining = max_chars - used - separator
        if remaining <= 0:
            truncated = True
            return True
        if len(value) >= remaining:
            parts.append(value[:remaining])
            used += separator + remaining
            truncated = True
            return True

        parts.append(value)
        used += separator + len(value)
        return False

    # iter_inner_content() yields Paragraph/Table blocks lazily in document
    # order. Avoid materializing a full paragraph list or collecting all tables
    # before applying the extraction limit.
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            if append_part(block.text):
                break
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                ]
                if append_part(" | ".join(cell for cell in cells if cell)):
                    break
            if truncated:
                break

    content = "\n".join(parts).strip()
    if truncated and content:
        return f"{content}\n\n[تم اختصار المحتوى بسبب الحد المسموح]"
    return _limit(content, max_chars)


def _read_xlsx_file(path: Path, max_chars: int) -> str:
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    parts: list[str] = []
    used = 0
    truncated = False

    def append_line(value: str) -> bool:
        nonlocal used, truncated
        separator = 1 if parts else 0
        remaining = max_chars - used - separator
        if remaining <= 0:
            truncated = True
            return True
        if len(value) >= remaining:
            parts.append(value[:remaining])
            used += separator + remaining
            truncated = True
            return True
        parts.append(value)
        used += separator + len(value)
        return False

    def render_row(row, limit: int) -> tuple[str, bool]:
        """Build only the row prefix that can fit into the remaining text budget."""
        values: list[str] = []
        used_in_row = 0
        has_value = False
        for value in row:
            if value is None:
                continue
            cell = str(value).strip()
            if not cell:
                continue
            separator = " | " if has_value else ""
            remaining = limit - used_in_row - len(separator)
            if remaining <= 0:
                return "".join(values), True
            if len(cell) >= remaining:
                values.append(separator + cell[:remaining])
                return "".join(values), True
            values.append(separator + cell)
            used_in_row += len(separator) + len(cell)
            has_value = True
        return "".join(values), False

    try:
        for worksheet in workbook.worksheets:
            if append_line(f"[ورقة: {worksheet.title}]"):
                break
            for row in worksheet.iter_rows(values_only=True):
                remaining = max_chars - used - (1 if parts else 0)
                if remaining <= 0:
                    truncated = True
                    break
                line, row_truncated = render_row(row, remaining)
                if not line:
                    if row_truncated:
                        truncated = True
                        break
                    continue
                if append_line(line) or row_truncated:
                    truncated = True
                    break
            if truncated:
                break

        content = "\n".join(parts).strip()
        if truncated and content:
            return f"{content}\n\n[تم اختصار المحتوى بسبب الحد المسموح]"
        return _limit(content, max_chars)
    finally:
        workbook.close()


def extract_text(
    path: str | Path,
    content_type: str,
    max_chars: int = DEFAULT_MAX_CHARS,
) -> str | None:
    """يرجع النص المستخرج، أو None للملفات غير النصية (مثل الصور)."""
    if content_type not in SUPPORTED_TEXT_TYPES:
        return None

    path = Path(path)
    try:
        if content_type == "text/plain":
            return _read_text_file(path, max_chars)
        if content_type in {"text/csv", "application/vnd.ms-excel"}:
            return _read_csv_file(path, max_chars)
        if content_type == "application/pdf":
            return _read_pdf_file(path, max_chars)
        if content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return _read_docx_file(path, max_chars)
        if content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            return _read_xlsx_file(path, max_chars)
    except Exception as exc:
        raise FileTextExtractionError(str(exc)) from exc

    return None
