"""اختبارات استخراج النص من الملفات."""
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from app.services.file_text_extractor import extract_text


def test_extract_text_file(tmp_path: Path):
    path = tmp_path / "notes.txt"
    path.write_text("Hello Linux\nPython and Cloud", encoding="utf-8")

    text = extract_text(path, "text/plain")

    assert "Hello Linux" in text
    assert "Python and Cloud" in text


def test_extract_csv_file(tmp_path: Path):
    path = tmp_path / "data.csv"
    path.write_text("name,age\nMohamed,25\nSara,30\n", encoding="utf-8")

    text = extract_text(path, "text/csv")

    assert "name | age" in text
    assert "Mohamed | 25" in text


def test_extract_docx_file(tmp_path: Path):
    path = tmp_path / "doc.docx"
    document = Document()
    document.add_paragraph("Linux administration basics")
    document.save(path)

    text = extract_text(
        path,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert "Linux administration basics" in text


def test_extract_xlsx_file(tmp_path: Path):
    path = tmp_path / "sheet.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(["name", "score"])
    worksheet.append(["Mohamed", 95])
    workbook.save(path)

    text = extract_text(
        path,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    assert "[ورقة: Data]" in text
    assert "Mohamed | 95" in text


def test_extract_pdf_file(tmp_path: Path):
    path = tmp_path / "empty.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as file:
        writer.write(file)

    assert extract_text(path, "application/pdf") == ""


def test_docx_extraction_stops_reading_blocks_after_character_limit(
    tmp_path: Path, monkeypatch
):
    from app.services import file_text_extractor

    class FakeParagraph:
        def __init__(self, text: str):
            self.text = text

    class FakeTable:
        pass

    class TrackingBlocks:
        def __init__(self):
            self.visited = 0

        def __iter__(self):
            self.visited += 1
            yield FakeParagraph("A" * 1000)
            self.visited += 1
            yield FakeParagraph("This later paragraph should not be visited.")

    blocks = TrackingBlocks()

    class FakeDocument:
        def iter_inner_content(self):
            return iter(blocks)

    monkeypatch.setattr(file_text_extractor, "Document", lambda _path: FakeDocument())
    monkeypatch.setattr(file_text_extractor, "Paragraph", FakeParagraph)
    monkeypatch.setattr(file_text_extractor, "Table", FakeTable)

    text = file_text_extractor._read_docx_file(tmp_path / "fake.docx", max_chars=20)

    assert text.startswith("A" * 20)
    assert "[تم اختصار المحتوى بسبب الحد المسموح]" in text
    assert blocks.visited == 1
    assert "later paragraph" not in text


def test_docx_extraction_preserves_paragraph_and_table_order(tmp_path: Path):
    path = tmp_path / "ordered.docx"
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Table content"
    document.add_paragraph("After table")
    document.save(path)

    text = extract_text(
        path,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert text.index("Before table") < text.index("Table content")
    assert text.index("Table content") < text.index("After table")


def test_pdf_extraction_stops_at_character_limit_and_skips_later_pages(
    tmp_path: Path, monkeypatch
):
    """Large text from a page must not be accumulated beyond the configured limit."""
    from app.services import file_text_extractor

    class FakePage:
        def __init__(self, content: str):
            self.content = content
            self.extract_calls = 0

        def extract_text(self):
            self.extract_calls += 1
            return self.content

    first_page = FakePage("short")
    oversized_page = FakePage("A" * 1000)
    later_page = FakePage("This page must never be visited.")

    class TrackingPages:
        def __init__(self):
            self.visited = 0

        def __iter__(self):
            for page in (first_page, oversized_page, later_page):
                self.visited += 1
                yield page

    pages = TrackingPages()

    class FakeReader:
        def __init__(self, _path):
            self.pages = pages

    monkeypatch.setattr(file_text_extractor, "PdfReader", FakeReader)

    text = file_text_extractor._read_pdf_file(tmp_path / "fake.pdf", max_chars=20)

    assert text.startswith("short\n\n" + "A" * 13)
    assert "[تم اختصار المحتوى بسبب الحد المسموح]" in text
    assert pages.visited == 2
    assert first_page.extract_calls == 1
    assert oversized_page.extract_calls == 1
    assert later_page.extract_calls == 0
    assert "This page must never be visited." not in text


def test_xlsx_extraction_bounds_accumulated_rows_and_closes_workbook(
    tmp_path: Path, monkeypatch
):
    """A single oversized spreadsheet cell must not be accumulated into context."""
    from app.services import file_text_extractor

    class FakeWorksheet:
        title = "Data"

        def __init__(self):
            self.rows_yielded = 0

        def iter_rows(self, values_only=True):
            self.rows_yielded += 1
            yield ("A" * 1000,)
            self.rows_yielded += 1
            yield ("This later row must not be visited.",)

    worksheet = FakeWorksheet()

    class FakeWorkbook:
        worksheets = [worksheet]

        def __init__(self):
            self.closed = False

        def close(self):
            self.closed = True

    workbook = FakeWorkbook()
    monkeypatch.setattr(
        file_text_extractor,
        "load_workbook",
        lambda **_kwargs: workbook,
    )

    original_limit = file_text_extractor._limit

    def assert_bounded_input(text: str, max_chars: int) -> str:
        assert len(text) <= max_chars
        return original_limit(text, max_chars)

    monkeypatch.setattr(file_text_extractor, "_limit", assert_bounded_input)

    text = file_text_extractor._read_xlsx_file(tmp_path / "fake.xlsx", max_chars=20)

    assert text.startswith("[ورقة: Data]\n" + "A" * 7)
    assert "[تم اختصار المحتوى بسبب الحد المسموح]" in text
    assert worksheet.rows_yielded == 1
    assert workbook.closed
    assert "This later row must not be visited." not in text


def test_docx_extraction_stops_when_a_block_exactly_fills_character_budget(
    tmp_path: Path, monkeypatch
):
    from app.services import file_text_extractor

    class FakeParagraph:
        def __init__(self, text: str):
            self.text = text

    class FakeTable:
        pass

    class TrackingBlocks:
        def __init__(self):
            self.visited = 0

        def __iter__(self):
            self.visited += 1
            yield FakeParagraph("A" * 20)
            self.visited += 1
            yield FakeParagraph("This later paragraph must not be visited.")

    blocks = TrackingBlocks()

    class FakeDocument:
        def iter_inner_content(self):
            return iter(blocks)

    monkeypatch.setattr(file_text_extractor, "Document", lambda _path: FakeDocument())
    monkeypatch.setattr(file_text_extractor, "Paragraph", FakeParagraph)
    monkeypatch.setattr(file_text_extractor, "Table", FakeTable)

    text = file_text_extractor._read_docx_file(tmp_path / "fake.docx", max_chars=20)

    assert text.startswith("A" * 20)
    assert "[تم اختصار المحتوى بسبب الحد المسموح]" in text
    assert blocks.visited == 1
    assert "later paragraph" not in text


def test_pdf_extraction_stops_when_a_page_exactly_fills_character_budget(
    tmp_path: Path, monkeypatch
):
    from app.services import file_text_extractor

    class FakePage:
        def __init__(self, content: str):
            self.content = content
            self.extract_calls = 0

        def extract_text(self):
            self.extract_calls += 1
            return self.content

    first_page = FakePage("A" * 20)
    later_page = FakePage("This page must not be extracted.")

    class TrackingPages:
        def __init__(self):
            self.visited = 0

        def __iter__(self):
            for page in (first_page, later_page):
                self.visited += 1
                yield page

    pages = TrackingPages()

    class FakeReader:
        def __init__(self, _path):
            self.pages = pages

    monkeypatch.setattr(file_text_extractor, "PdfReader", FakeReader)

    text = file_text_extractor._read_pdf_file(tmp_path / "fake.pdf", max_chars=20)

    assert text.startswith("A" * 20)
    assert "[تم اختصار المحتوى بسبب الحد المسموح]" in text
    assert pages.visited == 1
    assert first_page.extract_calls == 1
    assert later_page.extract_calls == 0


def test_xlsx_extraction_stops_when_a_row_exactly_fills_character_budget(
    tmp_path: Path, monkeypatch
):
    from app.services import file_text_extractor

    class FakeWorksheet:
        title = "Data"

        def __init__(self):
            self.rows_yielded = 0

        def iter_rows(self, values_only=True):
            self.rows_yielded += 1
            yield ("A" * 7,)
            self.rows_yielded += 1
            yield ("This later row must not be visited.",)

    worksheet = FakeWorksheet()

    class FakeWorkbook:
        worksheets = [worksheet]

        def __init__(self):
            self.closed = False

        def close(self):
            self.closed = True

    workbook = FakeWorkbook()
    monkeypatch.setattr(
        file_text_extractor,
        "load_workbook",
        lambda **_kwargs: workbook,
    )

    text = file_text_extractor._read_xlsx_file(tmp_path / "fake.xlsx", max_chars=20)

    assert text.startswith("[ورقة: Data]\n" + "A" * 7)
    assert "[تم اختصار المحتوى بسبب الحد المسموح]" in text
    assert worksheet.rows_yielded == 1
    assert workbook.closed
    assert "This later row must not be visited." not in text
